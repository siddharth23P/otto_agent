"""Documents otto makes: Markdown into PDF, Word, Excel or PowerPoint, in this process.

The research route converts its report with a script on the workspace's interpreter, which a host
with code execution off (the Android app) does not have -- so on a phone there was no way to make a
file, and a request for one went to the phone's own apps instead (2026-09-17: "create a pdf from
this doc" drove My Files and Google Docs for minutes). These writers need no subprocess: pure-Python
libraries (reportlab, python-docx, openpyxl, python-pptx; lxml and Pillow have Android builds).

One small Markdown reader feeds all four: headings, paragraphs, bullet and numbered lists, tables
and fenced code. Emphasis and links are reduced to their text.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

FORMATS = ("pdf", "docx", "xlsx", "pptx", "md")
#: Where make_document writes, inside the workspace.
OUT_DIR = "documents"

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_RULE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


@dataclass
class Block:
    kind: str                       # heading | paragraph | bullets | numbered | table | code
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


def plain(text: str) -> str:
    """Markdown emphasis, code spans and links reduced to their text."""
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"\1", text)
    text = re.sub(r"(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)", r"\1", text)
    text = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", text)
    return re.sub(r"`([^`]*)`", r"\1", text).strip()


def _cells(line: str) -> list[str]:
    inner = line.strip()
    inner = inner[1:] if inner.startswith("|") else inner
    inner = inner[:-1] if inner.endswith("|") else inner
    return [plain(c) for c in re.split(r"(?<!\\)\|", inner)]


def parse(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []
    fence: list[str] | None = None

    def flush() -> None:
        if para:
            blocks.append(Block("paragraph", plain(" ".join(para))))
            para.clear()

    lines = markdown.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if fence is None:
                flush()
                fence = []
            else:
                blocks.append(Block("code", "\n".join(fence)))
                fence = None
            i += 1
            continue
        if fence is not None:
            fence.append(line)
            i += 1
            continue
        s = line.rstrip()
        if not s.strip():
            flush()
        elif m := _HEADING.match(s):
            flush()
            blocks.append(Block("heading", plain(m.group(2)), level=len(m.group(1))))
        elif s.lstrip().startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                if not _RULE.match(lines[i]):
                    rows.append(_cells(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            blocks.append(Block("table", rows=[r + [""] * (width - len(r)) for r in rows]))
            continue
        elif (m := _BULLET.match(s)) or (n := _NUMBERED.match(s)):
            flush()
            kind = "bullets" if m else "numbered"
            item = plain((m or n).group(1))
            if blocks and blocks[-1].kind == kind and not blocks[-1].text:
                blocks[-1].items.append(item)
            else:
                blocks.append(Block(kind, items=[item]))
        else:
            if blocks and blocks[-1].kind in ("bullets", "numbered") and not para and line.startswith("  "):
                blocks[-1].items[-1] += " " + plain(s)
            else:
                para.append(s.strip())
        i += 1
    if fence is not None:
        blocks.append(Block("code", "\n".join(fence)))
    flush()
    return blocks


def title_of(blocks: list[Block], fallback: str) -> str:
    return next((b.text for b in blocks if b.kind == "heading"), fallback)


# -- writers -----------------------------------------------------------------------------------

def _docx(blocks: list[Block], target: Path, title: str) -> None:
    from docx import Document

    doc = Document()
    doc.core_properties.title = title
    for b in blocks:
        if b.kind == "heading":
            doc.add_heading(b.text, min(b.level - 1, 4) if b.level > 1 else 0)
        elif b.kind == "paragraph":
            doc.add_paragraph(b.text)
        elif b.kind in ("bullets", "numbered"):
            style = "List Bullet" if b.kind == "bullets" else "List Number"
            for item in b.items:
                doc.add_paragraph(item, style=style)
        elif b.kind == "table":
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = value
                    if r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
        elif b.kind == "code":
            for code_line in b.text.split("\n"):
                doc.add_paragraph(code_line, style="No Spacing")
    doc.save(str(target))


def _pdf(blocks: list[Block], target: Path, title: str) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    heading = {1: styles["Title"], 2: styles["Heading1"], 3: styles["Heading2"]}
    body = styles["BodyText"]
    story = []
    for b in blocks:
        if b.kind == "heading":
            story.append(Paragraph(escape(b.text), heading.get(b.level, styles["Heading3"])))
        elif b.kind == "paragraph":
            story += [Paragraph(escape(b.text), body), Spacer(1, 0.2 * cm)]
        elif b.kind in ("bullets", "numbered"):
            story.append(ListFlowable([ListItem(Paragraph(escape(i), body)) for i in b.items],
                                      bulletType="bullet" if b.kind == "bullets" else "1"))
            story.append(Spacer(1, 0.2 * cm))
        elif b.kind == "table":
            data = [[Paragraph(escape(v), body) for v in row] for row in b.rows]
            table = Table(data, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story += [table, Spacer(1, 0.3 * cm)]
        elif b.kind == "code":
            story.append(Preformatted(b.text, styles["Code"]))
    SimpleDocTemplate(str(target), pagesize=A4, title=title, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm).build(story or [Paragraph("", body)])


def _number(value: str):
    """A cell's value as a number when it reads as one ("1,200", "12.5", "40%"), else the text."""
    s = value.strip().replace(",", "")
    try:
        if s.endswith("%"):
            return float(s[:-1]) / 100
        return int(s) if re.fullmatch(r"-?\d+", s) else float(s)
    except ValueError:
        return value


def _xlsx(blocks: list[Block], target: Path, title: str) -> None:
    """Every table on a sheet of its own (named by the heading above it), numbers as numbers; any
    other text on a "Notes" sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()

    def sheet(name: str):
        base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:28] or "Sheet"
        unique, n = base, 2
        while unique.lower() in used:
            unique, n = f"{base[:25]}-{n}", n + 1
        used.add(unique.lower())
        return wb.create_sheet(unique)

    last_heading = title
    notes = []
    for b in blocks:
        if b.kind == "heading":
            last_heading = b.text
            notes.append([b.text])
        elif b.kind == "table":
            ws = sheet(last_heading)
            for r, row in enumerate(b.rows):
                ws.append(row if r == 0 else [_number(v) for v in row])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            for c in range(1, len(b.rows[0]) + 1):
                width = max(len(str(ws.cell(r, c).value or "")) for r in range(1, ws.max_row + 1))
                ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 60)
        elif b.kind == "paragraph" or b.kind == "code":
            notes.append([b.text])
        else:
            notes += [[("• " if b.kind == "bullets" else f"{k}. ") + item] for k, item in enumerate(b.items, 1)]
    has_text = any(b.kind not in ("heading", "table") for b in blocks)
    if has_text or (notes and not wb.sheetnames):
        ws = sheet("Notes")
        for row in notes:
            ws.append(row)
        ws.column_dimensions["A"].width = 100
    if not wb.sheetnames:
        wb.create_sheet("Sheet")
    wb.save(str(target))


def _pptx(blocks: list[Block], target: Path, title: str) -> None:
    """A title slide, then a slide per heading: its paragraphs and list items as bullets (at most
    eight a slide, the rest on continuation slides), and a table as a table."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    subtitle = next((b.text for b in blocks if b.kind == "paragraph"), "")
    if len(first.placeholders) > 1:
        first.placeholders[1].text = subtitle[:200]

    slides: list[tuple[str, list]] = []
    current: tuple[str, list] | None = None
    for b in blocks:
        if b.kind == "heading":
            if b.text == title and current is None:
                continue
            current = (b.text, [])
            slides.append(current)
            continue
        if current is None:
            current = (title, [])
            slides.append(current)
        if b.kind == "paragraph" and b.text != subtitle:
            current[1].append(b.text)
        elif b.kind in ("bullets", "numbered"):
            current[1].extend(b.items)
        elif b.kind == "table":
            current[1].append(b)
        elif b.kind == "code":
            current[1].extend(b.text.split("\n")[:8])

    for heading, items in slides:
        bullets = [i for i in items if isinstance(i, str)]
        tables = [i for i in items if isinstance(i, Block)]
        chunks = [bullets[k:k + 8] for k in range(0, len(bullets), 8)] or ([[]] if not tables else [])
        for n, chunk in enumerate(chunks):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = heading if n == 0 else f"{heading} (cont.)"
            frame = slide.placeholders[1].text_frame
            for k, text in enumerate(chunk):
                para = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
                para.text = text[:300]
                para.font.size = Pt(20 if len(chunk) <= 5 else 16)
        for table_block in tables:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = heading
            rows, cols = len(table_block.rows), len(table_block.rows[0])
            shape = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.4) * rows)
            for r, row in enumerate(table_block.rows):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
                    shape.table.cell(r, c).text_frame.paragraphs[0].font.size = Pt(14)
    prs.save(str(target))


WRITERS = {"docx": _docx, "pdf": _pdf, "xlsx": _xlsx, "pptx": _pptx}


def write(markdown: str, target: Path, fmt: str, title: str = "") -> Path:
    """`markdown` as a `fmt` file at `target`. Raises ValueError for an unknown format, and the
    writer's own error (a missing library, a bad table) otherwise."""
    if fmt not in FORMATS:
        raise ValueError(f"no {fmt!r} format; choose one of {', '.join(FORMATS)}")
    target.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "md":
        target.write_text(markdown, encoding="utf-8")
        return target
    blocks = parse(markdown)
    WRITERS[fmt](blocks, target, title or title_of(blocks, target.stem))
    return target
